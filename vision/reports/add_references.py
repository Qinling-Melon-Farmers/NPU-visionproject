"""Insert in-text references for tables and figures that lack one."""
from __future__ import annotations

import re

from docx.shared import Pt
from docx.text.paragraph import Paragraph

from docx_helpers import _set_run_fonts


def add_missing_references(doc):
    existing = set()
    for paragraph in doc.paragraphs:
        if paragraph.style.name != "Normal":
            continue
        existing.update(int(x) for x in re.findall(r"表\s*(\d+)", paragraph.text))
    table_refs = set(existing)
    figure_refs = set()
    for paragraph in doc.paragraphs:
        if paragraph.style.name != "Normal":
            continue
        figure_refs.update(int(x) for x in re.findall(r"图\s*(\d+)", paragraph.text))

    for paragraph in list(doc.paragraphs):
        if paragraph.style.name != "Caption":
            continue
        match = re.match(r"表\s*(\d+)\s+(.*)", paragraph.text)
        if match:
            number = int(match.group(1))
            if number in table_refs:
                continue
            lead = paragraph.insert_paragraph_before(f"表 {number} 汇总了{match.group(2)}。")
            lead.style = doc.styles["Normal"]
            lead.paragraph_format.first_line_indent = Pt(24)
            lead.paragraph_format.line_spacing = 1.5
            for run in lead.runs:
                _set_run_fonts(run, "宋体", 12)
            table_refs.add(number)
            continue
        match = re.match(r"图\s*(\d+)\s+(.*)", paragraph.text)
        if match:
            number = int(match.group(1))
            if number in figure_refs:
                continue
            previous = paragraph._p.getprevious()
            target = Paragraph(previous, paragraph._parent) if previous is not None and previous.tag.endswith("}p") else paragraph
            lead = target.insert_paragraph_before(f"图 {number} 展示了{match.group(2)}。")
            lead.style = doc.styles["Normal"]
            lead.paragraph_format.first_line_indent = Pt(24)
            lead.paragraph_format.line_spacing = 1.5
            for run in lead.runs:
                _set_run_fonts(run, "宋体", 12)
            figure_refs.add(number)
