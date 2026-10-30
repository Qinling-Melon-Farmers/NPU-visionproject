"""Build the revised multimedia vision report.

Run with the project Conda environment:
    python -s vision/reports/build_report_v2.py
"""
from __future__ import annotations

import json
import subprocess
from pathlib import Path

from docx.shared import Pt

from docx_helpers import (
    add_field,
    add_footer_page_x_of_y,
    add_footer_roman,
    add_header,
    cover_info_table,
    cover_line,
    enable_field_update,
    inspect,
    new_section,
    set_core_properties,
    set_page_number_format,
    setup_document,
)
from report_appendices import (
    add_appendix_a,
    add_appendix_b,
    add_appendix_c,
    add_appendix_d,
)
from report_sections import (
    add_abstract,
    add_conclusion,
    add_discussion,
    add_environment,
    add_image,
    add_intro,
    add_references,
    add_summary,
    add_video,
)
from add_references import add_missing_references

ROOT = Path(__file__).resolve().parents[2]
VISION = ROOT / "vision"
RUN = VISION / "outputs" / "visual-20261023-fp32"
REPORT_DIR = VISION / "reports"
OUT = REPORT_DIR / "delivery-20261028" / "vision_report.docx"


def load_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8-sig"))


def load_jsonl(path: Path):
    return [json.loads(line) for line in path.read_text(encoding="utf-8-sig").splitlines() if line.strip()]


def git_commit() -> str:
    try:
        return subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=ROOT, text=True).strip()
    except Exception:
        return "unknown"


def collect_run_meta():
    items = []
    for kind in ("image", "video", "summary"):
        for path in sorted(RUN.glob(f"{kind}_run_*.json")):
            data = load_json(path)
            command = data.get("command", [])
            if len(command) > 1 and command[1] == kind:
                items.append(data)
                break
    reports = []
    for path in sorted(RUN.glob("report_run_*.json")):
        data = load_json(path)
        command = data.get("command", [])
        if len(command) > 1 and command[1] == "report":
            reports.append((path.stat().st_mtime, data))
    if reports:
        items.append(max(reports, key=lambda item: item[0])[1])
    return items


def build():
    image = load_json(RUN / "image" / "results.json")
    video = load_json(RUN / "video" / "results.json")
    summary = load_json(RUN / "summary" / "results.json")
    env = load_json(VISION / "manifests" / "environment_validation.json")
    audit_pointer = load_json(VISION / "logs" / "interaction_audit_latest.json")
    audit = load_json(ROOT / audit_pointer["audit_path"])
    records = load_jsonl(ROOT / audit_pointer["transcript_path"])
    student_records = [r for r in records if r.get("category") == "student_main"]
    agent_records = [r for r in records if r.get("category") == "agent_collaboration"]
    corrections = load_jsonl(VISION / "logs" / "implementation_corrections.jsonl")

    ctx = {
        "run": RUN,
        "vision": VISION,
        "reports": REPORT_DIR,
        "figures": REPORT_DIR / "figures_v2",
        "image": image,
        "video": video,
        "summary": summary,
        "env": env,
        "audit": audit,
        "student_records": student_records,
        "agent_records": agent_records,
        "corrections": corrections,
        "run_meta": collect_run_meta(),
        "commit": git_commit(),
    }

    doc = setup_document()
    set_core_properties(doc)

    # cover section
    cover = doc.sections[0]
    cover.footer.is_linked_to_previous = False
    cover.header.is_linked_to_previous = False
    for _ in range(3):
        cover_line(doc, "", 12)
    cover_line(doc, "西北工业大学", 22, True, "方正舒体简体")
    cover_line(doc, "《多媒体技术》实验报告", 22, True)
    cover_line(doc, "", 12)
    cover_line(doc, "视觉部分", 14)
    cover_line(doc, "GeM 图像检索 · ViSiL 视频检索 · DSNet 无锚框视频摘要", 13)
    for _ in range(2):
        cover_line(doc, "", 12)
    cover_info_table(doc, [
        ("学    院：", "计算机学院"),
        ("学    号：", "2023302610"),
        ("姓    名：", "肖怀瑾"),
        ("班    级：", "10032301"),
        ("专    业：", "数据科学与大数据技术"),
        ("实验时间：", "2026 年 10 月 23 日"),
    ])
    for _ in range(3):
        cover_line(doc, "", 12)
    cover_line(doc, "2026 年 10 月", 12, True)

    # front matter
    front = new_section(doc)
    add_header(front, "西北工业大学《多媒体技术》实验报告")
    add_footer_roman(front)
    set_page_number_format(front, "upperRoman", 1)
    add_abstract(doc, ctx)
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == "摘要":
            paragraph.paragraph_format.page_break_before = False
            break
    doc.add_page_break()
    toc_title = doc.add_paragraph()
    toc_title.alignment = 1
    toc_title.paragraph_format.space_before = Pt(0)
    toc_title.paragraph_format.space_after = Pt(12)
    run = toc_title.add_run("目录")
    from docx_helpers import _set_run_fonts

    _set_run_fonts(run, "楷体", 16, True, None)
    toc = doc.add_paragraph()
    toc.paragraph_format.first_line_indent = Pt(0)
    add_field(toc, 'TOC \\o "1-2" \\h \\z \\u', "右键更新目录")

    # body
    body_section = new_section(doc)
    add_header(body_section, "西北工业大学《多媒体技术》实验报告")
    add_footer_page_x_of_y(body_section)
    set_page_number_format(body_section, None, 1)
    add_intro(doc, ctx)
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == "1 引言":
            paragraph.paragraph_format.page_break_before = False
            break
    add_environment(doc, ctx)
    add_image(doc, ctx)
    add_video(doc, ctx)
    add_summary(doc, ctx)
    add_discussion(doc, ctx)
    add_conclusion(doc, ctx)
    add_references(doc, ctx)
    add_appendix_a(doc, ctx)
    add_appendix_b(doc, ctx)
    add_appendix_c(doc, ctx)
    add_appendix_d(doc, ctx)
    add_missing_references(doc)

    enable_field_update(doc)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    if OUT.exists():
        raise FileExistsError(OUT)
    doc.save(OUT)
    inspect(doc, OUT)


if __name__ == "__main__":
    build()
