"""Appendices for the multimedia vision report."""
from __future__ import annotations

import csv
import json
import re
from pathlib import Path

from docx.shared import Pt

from docx_helpers import (
    add_box,
    add_code_block,
    add_figure,
    add_pageref,
    add_table,
    body,
    heading,
    _cant_split,
    _cell_margins,
    _repeat_header,
    _set_cell_bottom_border,
    _set_run_fonts,
    _set_table_borders,
    _set_table_width,
    _shade,
    clean_latex,
)

T_APER = 17
T_ACTRL = 18
T_ABASE = 19
T_AFILES = 20
T_RUNS = 21
T_CORR = 22
T_FAIL = 23
T_CSTAT = 24
T_QA = 25
T_DSTAT = 26
T_DTABLE = 27
T_DGAP = 28


def _short(text: str, limit: int = 46) -> str:
    text = clean_latex(text)
    text = re.sub(r"# Context from my IDE setup:.*?My request:", "", text, flags=re.S)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"[*#`]", "", text)
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return "（空）"
    first = re.split(r"[。\n]", text)[0].strip()
    return first[:limit] + ("…" if len(first) > limit else "")


def _meta_line(doc, items):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.keep_with_next = True
    text = "  ·  ".join(f"{label}：{value}" for label, value in items)
    run = paragraph.add_run(text)
    _set_run_fonts(run, "宋体", 7.5, False, None)


def _time(record) -> str:
    value = str(record.get("timestamp", ""))
    match = re.match(r"(\d{4}-\d{2}-\d{2})T(\d{2}:\d{2}:\d{2})", value)
    return f"{match.group(1)} {match.group(2)}" if match else value


def _kind_zh(kind: str) -> str:
    return {
        "visible_message": "消息",
        "structured_question": "结构化问",
        "structured_answer": "结构化答",
        "coordination_spawn_agent": "创建代理",
        "coordination_followup_task": "追加任务",
        "coordination_send_message": "代理消息",
        "delivered_agent_message": "代理回复",
    }.get(kind, kind)


def _phase(index: int) -> str:
    if index <= 5:
        return "材料"
    if index <= 15:
        return "规划"
    if index <= 27:
        return "计划"
    if index <= 42:
        return "运行"
    return "归档"


def _parse_questions(text: str):
    try:
        return json.loads(text).get("questions", [])
    except Exception:
        return []


def _parse_answers(text: str):
    try:
        return json.loads(text).get("answers", {})
    except Exception:
        return {}


def _answer_text(answer) -> str:
    if isinstance(answer, dict):
        values = answer.get("answers", [])
        return "；".join(str(v) for v in values)
    return str(answer)


def _structured_display(record):
    kind = record.get("kind")
    text = record.get("text", "")
    if kind == "structured_question":
        lines = []
        for question in _parse_questions(text):
            options = " / ".join(item.get("label", "") for item in question.get("options", []))
            lines.append(f"问题：{question.get('question', '')}")
            lines.append(f"选项：{options}")
        return lines or ["（结构化问题原文见 C.6）"]
    if kind == "structured_answer":
        answers = _parse_answers(text)
        return [f"{key}：{_answer_text(value)}" for key, value in answers.items()] or ["（结构化回答原文见 C.6）"]
    return None


def _topic(record, limit: int = 40) -> str:
    kind = record.get("kind")
    text = record.get("text", "")
    if kind == "structured_question":
        names = [item.get("header") or item.get("id", "") for item in _parse_questions(text)]
        return "结构化提问：" + "、".join(names)
    if kind == "structured_answer":
        labels = {
            "data_source": "数据来源",
            "deadline": "截止时间",
            "experiment_depth": "实验深度",
            "research_focus": "研究重点",
            "asset_location": "资源位置",
        }
        names = [labels.get(key, key) for key in _parse_answers(text)]
        return "结构化回答：" + "、".join(names)
    return _short(text, limit)


def _index_table(doc, records, bookmarks):
    headers = ["编号", "时间", "方向", "阶段", "主题", "页"]
    widths = [1.9, 2.5, 2.0, 1.8, 4.6, 1.8]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Normal Table"
    table.autofit = False
    table.alignment = 1
    _set_table_borders(table)
    _set_table_width(table, widths)
    _cell_margins(table, top=20, bottom=20, left=40, right=40)
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.alignment = 1
        run = paragraph.add_run(value)
        _set_run_fonts(run, "宋体", 8, True, None)
        _shade(cell, "F2F2F2")
        _set_cell_bottom_border(cell, 6)
    _repeat_header(table.rows[0])
    for index, record in enumerate(records, 1):
        row = table.add_row()
        _cant_split(row)
        values = [
            f"C-{index:02d}", _time(record)[11:19],
            "学→模" if record.get("sender") == "student" else "模→学",
            _phase(index), _topic(record, 32),
        ]
        for cell, value in zip(row.cells[:5], values):
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            _set_run_fonts(run, "宋体", 7.5, False, None)
        page_cell = row.cells[5]
        page_cell.text = ""
        paragraph = page_cell.paragraphs[0]
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.alignment = 1
        add_pageref(paragraph, bookmarks[index - 1])
        for run in paragraph.runs:
            _set_run_fonts(run, "宋体", 7.5, False, None)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(4)


def add_appendix_a(doc, ctx):
    video = ctx["video"]
    heading(doc, "附录 A 逐查询结果与补充图表", 1)
    body(
        doc,
        "本附录保存主方法的逐查询指标、视频对照的逐查询结果、全部 Top-3 代表帧、随机基线和结果文件索引。"
        "表格直接读取本次运行的 per_query.csv、top5.csv、results.json 和随机基线记录，不重算、不改写分数。"
        "完整排名和分数矩阵保存在 vision/outputs/visual-20261023-fp32 的 NPZ 文件中。",
    )
    per = list(csv.DictReader((ctx["run"] / "image" / "per_query.csv").open(encoding="utf-8-sig", newline="")))
    selected = [r for r in per if r["gallery"] == "full" and r["method"] == "gem"]
    groups: dict[str, dict[int, float]] = {}
    for row in selected:
        match = re.search(r"__(\d+)", row["query_id"])
        if match:
            groups.setdefault(row["object_id"], {})[int(match.group(1))] = float(row["p_at_5"])
    rows = []
    for object_id in sorted(groups, key=int):
        values = [groups[object_id].get(angle) for angle in (0, 90, 180, 270)]
        mean = sum(v for v in values if v is not None) / 4
        rows.append([object_id] + [f"{v:.2f}" if v is not None else "—" for v in values] + [f"{mean:.2f}"])
    add_table(
        doc, T_APER, "主图库 GeM 逐查询 P@5（按对象×视角，100 行）",
        ["对象 ID", "0°", "90°", "180°", "270°", "平均"],
        rows, widths=[2.6, 2.4, 2.4, 2.4, 2.4, 2.4], font_size=8,
        numeric_cols={1, 2, 3, 4, 5}, align_center_cols={0}, keep_rows=True,
    )

    control_rows = []
    for item in video["control_queries"]:
        top3 = "\n".join(
            f"{c['id']}（{c['score']:.4f}，{'相关' if c['relevant'] else '不相关'}）"
            for c in item["top3"]
        )
        control_rows.append([item["query_id"], item["source_id"], str(item["hit_at_3"]), f"{item['ap']:.4f}", top3])
    add_table(
        doc, T_ACTRL, "去 CNN 对照的逐查询结果",
        ["查询 ID", "源视频", "Hit@3", "AP", "Top-3 候选（相似度，相关性）"],
        control_rows, widths=[4.0, 3.0, 1.3, 1.4, 4.9], font_size=7.5,
        numeric_cols={2, 3}, align_center_cols={1},
    )

    body(
        doc,
        "图 15 至图 32 是正文没有展开的其余 18 个视频查询的 Top-3 代表帧，按查询顺序排列。"
        "每张图的标题给出查询和候选的完整 ID，分数是完整 ViSiL 的实际输出；同一查询的候选顺序按分数降序。",
    )
    figure_no = 15
    for path in video["top3_figures"][2:]:
        query_id = Path(path).stem.replace("_query_top3", "")
        add_figure(doc, figure_no, path, f"查询 {query_id} 的 Top-3 代表帧", width_cm=14.5)
        figure_no += 1

    baseline = ctx["summary"]["random_baseline"]["runs"]
    rows = [[str(item["seed"]), f"{item['avg_f1']:.4f}", f"{item['ratio'] * 100:.4f}%"] for item in baseline]
    rows.append(["均值", f"{ctx['summary']['random_baseline']['mean_f1']:.4f}",
                 f"{sum(i['ratio'] for i in baseline) / len(baseline) * 100:.4f}%"])
    add_table(
        doc, T_ABASE, "15% 预算下随机片段基线的 30 次运行",
        ["种子", "平均 F1", "实际比例"],
        rows, widths=[3.0, 5.8, 5.8], font_size=8.5, numeric_cols={1, 2}, align_center_cols={0},
    )

    add_table(
        doc, T_AFILES, "结果文件与可复核内容",
        ["文件", "内容", "用途"],
        [
            ["image/per_query.csv", "全部方法、图库、查询的 P@5 与 AP", "逐查询核对"],
            ["image/top5.csv", "全部 Top-5 标识、真实相似度、相关标签", "案例复核"],
            ["image/rankings_*.npz", "完整候选排名与分数矩阵", "重算指标"],
            ["video/results.json", "主方法与对照的评分、Top-3、命中、AP", "视频检索复核"],
            ["video/relevance_labels.csv", "20×90 查询—候选相关性标签", "标签复核"],
            ["summary/prediction.npz", "重要性、中心度、置信度、偏移、候选区间、NMS 结果", "摘要复核"],
            ["summary/summary_*.npz", "三种预算的选帧掩码、分段、恢复分数", "选段复核"],
        ],
        widths=[4.4, 7.4, 2.8], font_size=8.5,
    )


def add_appendix_b(doc, ctx):
    heading(doc, "附录 B 运行、修正与复现记录", 1)
    body(
        doc,
        "本附录只保留与实验结论有关的运行信息和修正记录，不把完整 JSON 原文整段复制进报告。"
        "完整运行元数据、日志和修正 JSONL 保存在 vision/outputs/visual-20261023-fp32 与 vision/logs，"
        "文件 SHA256 可在提交包清单和运行元数据中核对。",
    )
    run_zh = {"image": "图像", "video": "视频", "summary": "摘要", "report": "报告", "verify": "验收"}
    status_zh = {"completed": "完成", "passed": "通过", "failed": "失败"}
    rows = []
    for item in ctx["run_meta"]:
        command = item.get("command", [])
        kind = command[1] if len(command) > 1 else "run"
        rows.append([
            run_zh.get(kind, kind),
            f"python -m vision {kind} --run-id visual-20261023-fp32",
            status_zh.get(item.get("status", ""), item.get("status", "")),
            f"{item.get('elapsed_seconds', 0):.2f}" if item.get("elapsed_seconds") is not None else "—",
            str(item.get("code", {}).get("commit", ""))[:12],
        ])
    rows.append(["验收", "python -m vision verify --run-id visual-20261023-fp32", "通过", "—", ctx["commit"][:12]])
    add_table(
        doc, T_RUNS, "正式运行与验收记录",
        ["运行", "命令", "状态", "耗时（秒）", "代码版本"],
        rows, widths=[1.8, 7.0, 2.0, 1.8, 2.0], font_size=7.5,
        numeric_cols={3}, align_center_cols={0, 2, 4},
    )

    stage_zh = {
        "environment": "环境",
        "video preparation": "视频准备",
        "summary": "摘要",
        "reproducibility": "可复现",
        "report": "报告",
        "playback": "播放",
        "final verification and layout": "验收与排版",
        "interaction provenance and package": "交互归档与打包",
    }
    correction_rows = [
        [stage_zh.get(item.get("stage", ""), item.get("stage", "")), item.get("issue", "—"),
         item.get("action", ""), item.get("outcome", "")]
        for item in ctx["corrections"]
    ]
    add_table(
        doc, T_CORR, "实施过程中的主要修正",
        ["阶段", "问题", "修正", "结果"],
        correction_rows, widths=[2.2, 4.2, 5.2, 3.0], font_size=7.5,
    )

    add_table(
        doc, T_FAIL, "失败、重跑与验收证据",
        ["问题", "处理", "证据"],
        [
            ["首轮 FP32 重放分数差 4.798×10⁻⁵，超过 1×10⁻⁶ 容差", "关闭 TF32/benchmark，启用确定性卷积，写入缓存指纹，新 run-id 重跑", "replay_image.json；最大差 4.1723×10⁻⁷"],
            ["Conda 缓存元数据缺 URL，离线安装 AssertionError", "备份缓存元数据，按已核验 URL/SHA256 补全，@EXPLICIT 安装", "logs/conda-create-*.json / stderr"],
            ["Conda FFmpeg 无 libx264，OpenCV 无 Video I/O", "固定 MPEG-4 编码；用环境内 FFmpeg 管道解码；保持作者采样", "video_sources.json；合成帧序测试"],
            ["TVSum H5 无 video_name，首次读取 KeyError", "用帧数、标注曲线、原视频哈希建立映射；不修改 H5", "tvsum_h5_mapping.json；video_1 10,597 帧"],
            ["报告生成器字段不一致、Word 自动化连到 WPS", "统一播放字段；显式 Word.Application.16；保留失败输出", "implementation_corrections.jsonl；页面预览"],
        ],
        widths=[3.8, 6.4, 4.4], font_size=8,
    )

    body(doc, "复现时从工作区根目录运行下列命令。正式结果使用 run-id visual-20261023-fp32；重跑必须换新的 run-id，不能覆盖已有证据。")
    add_code_block(
        doc,
        "\n".join([
            "conda run --no-capture-output -p E:\\MultimediaCourseAssets\\envs\\multimedia-vision python -s -m vision image --run-id visual-20261023-fp32",
            "conda run --no-capture-output -p E:\\MultimediaCourseAssets\\envs\\multimedia-vision python -s -m vision video --run-id visual-20261023-fp32",
            "conda run --no-capture-output -p E:\\MultimediaCourseAssets\\envs\\multimedia-vision python -s -m vision summary --run-id visual-20261023-fp32",
            "conda run -p E:\\MultimediaCourseAssets\\envs\\multimedia-vision python -s -m pytest vision/tests",
            "conda run -p E:\\MultimediaCourseAssets\\envs\\multimedia-vision python -s vision/scripts/replay_image.py --config vision/configs/default.yaml --run-root vision/outputs/visual-20261023-fp32",
            "conda run -p E:\\MultimediaCourseAssets\\envs\\multimedia-vision python -s vision/scripts/replay_models.py --kind video --config vision/configs/default.yaml --run-root vision/outputs/visual-20261023-fp32",
            "conda run -p E:\\MultimediaCourseAssets\\envs\\multimedia-vision python -s vision/scripts/replay_models.py --kind summary --config vision/configs/default.yaml --run-root vision/outputs/visual-20261023-fp32",
            "conda run -p E:\\MultimediaCourseAssets\\envs\\multimedia-vision python -s -m vision verify --run-id visual-20261023-fp32",
        ]),
    )
    body(
        doc,
        "提交包只包含代码、配置、测试、说明、来源清单、许可证、报告和必要的运行证据，不含数据、权重、特征缓存或摘要 MP4。"
        "选择性打包脚本是 vision/scripts/package_submission.py，解压校验脚本是 vision/scripts/verify_submission.py。",
    )


def add_appendix_c(doc, ctx):
    records = ctx["student_records"]
    heading(doc, "附录 C 学生与编程大模型的交互记录", 1)
    body(
        doc,
        "本附录收录学生直接输入、编程大模型的可见回复和两批结构化问答，不包含代理之间的任务分派、隐藏推理或内部指令。"
        "原文来自 interactions_20261023_audit02.jsonl，字符不改写；为便于阅读，IDE 自动附加上下文折叠，Markdown 标记渲染为格式，"
        "结构化问答的 JSON 转为表格，完整原文仍保存在日志中。2026-10-28 的交付审阅不属于实验执行交互，未写入本附录。",
    )
    counts = ctx["audit"]["counts"]
    questions = []
    for record in records:
        if record.get("kind") == "structured_question":
            questions.extend(_parse_questions(record.get("text", "")))
    add_table(
        doc, T_CSTAT, "交互记录统计",
        ["口径", "数量"],
        [
            ["学生自由文本提示", str(counts["student_free_text_prompts"])],
            ["结构化问答", f"{counts['structured_question_batches']} 批 / {len(questions)} 问"],
            ["学生输入合计", str(counts["student_free_text_prompts"] + len(questions))],
            ["编程大模型可见回复", str(counts["main_agent_visible_replies"])],
            ["学生—编程大模型记录", str(len(records))],
            ["代理协作记录（另列）", f"{counts['agent_coordination_records']}（17 可读 / 62 加密）"],
            ["会话跨度", "2026-10-23 06:22:27–08:04:10（1 小时 41 分 43 秒；含下载与并发，不作为劳动时长）"],
            ["外部等待", "未单独计时，报告不编造精确用时"],
        ],
        widths=[5.6, 9.0], font_size=8.5,
    )
    body(
        doc,
        "索引按原快照顺序排列，编号 C-01 至 C-46 与日志行号一一对应；每条记录前的小标题是编辑归纳的主题，正文仍是原始消息。"
        "页码在 Word 更新域后生成。",
    )
    bookmarks = [f"C{index:02d}" for index in range(1, len(records) + 1)]
    _index_table(doc, records, bookmarks)

    phase_names = {
        1: "阶段 I：材料与规则",
        6: "阶段 II：规划与结构化问答",
        16: "阶段 III：计划、目录与 Git",
        28: "阶段 IV：环境、实验与报告",
        43: "阶段 V：交互归档",
    }
    from docx_helpers import add_bookmark

    for index, (record, bookmark) in enumerate(zip(records, bookmarks), 1):
        if index in phase_names:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.space_before = Pt(8)
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.keep_with_next = True
            run = paragraph.add_run(phase_names[index])
            _set_run_fonts(run, "黑体", 10.5, True, None)
        title = doc.add_paragraph(style="Transcript Title")
        run = title.add_run(f"C-{index:02d}  {_topic(record, 34)}")
        _set_run_fonts(run, "宋体", 10.5, True, None)
        add_bookmark(title, bookmark)
        sender = "学生" if record.get("sender") == "student" else "编程大模型"
        model = record.get("model", "")
        if model == "not_applicable_human":
            model = "不适用"
        _meta_line(doc, [
            ("时间", _time(record)), ("方向", f"{sender}→{'模型' if sender == '学生' else '学生'}"),
            ("类型", _kind_zh(record.get("kind"))), ("模型", model),
            ("来源", f"{Path(str(record.get('source_file', ''))).name}:{record.get('source_line', '')}"),
            ("记录", str(record.get("id", ""))[:8] + "…"), ("正文", "完整" if record.get("body_available") else "加密"),
        ])
        structured = _structured_display(record)
        if structured is not None:
            add_box(doc, structured, fill="FFFFFF", border_color="BFBFBF", font_size=8.5,
                    label=sender, keep_together=False)
            note = doc.add_paragraph()
            note.paragraph_format.first_line_indent = Pt(0)
            note.paragraph_format.space_after = Pt(4)
            run = note.add_run("原文（JSON）见 C.6；此处只列问题、选项和实际选择。")
            _set_run_fonts(run, "宋体", 8, False, None)
            continue
        text = clean_latex(record.get("text", ""))
        if record.get("sender") == "student":
            match = re.search(r"##\s*My request:\s*(.*)", text, flags=re.S)
            if match:
                text = "（IDE 自动附加上下文已折叠）\n" + match.group(1).strip()
        keep = len(text) < 1200
        add_box(doc, text.split("\n"), fill="F3F4F6" if sender == "学生" else "FFFFFF",
                border_color="7F7F7F" if sender == "学生" else "BFBFBF",
                font_size=8.5, label=sender, keep_together=keep)
        if index in (9, 12, 16, 19, 35, 36, 43, 46):
            note = doc.add_paragraph()
            note.paragraph_format.first_line_indent = Pt(0)
            note.paragraph_format.space_after = Pt(4)
            run = note.add_run("编者注（非原文）：" + _editor_note(index))
            _set_run_fonts(run, "宋体", 8, False, None)

    qa_heading = heading(doc, "C.6 结构化问答", 2)
    qa_heading.paragraph_format.page_break_before = True
    body(doc, "两批结构化问答的问题、选项和实际选择如下。选项中的“Recommended”是当时界面给出的推荐标记，不是学生的补充说明。")
    answer_map = {}
    for record in records:
        if record.get("kind") == "structured_answer":
            answer_map.update(_parse_answers(record.get("text", "")))
    qa_rows = []
    for question in questions:
        qid = question.get("id", "")
        options = " / ".join(item.get("label", "") for item in question.get("options", []))
        answer = answer_map.get(qid, {})
        qa_rows.append([question.get("question", ""), options,
                        _answer_text(answer) if answer else "未记录", question.get("header", "")])
    add_table(
        doc, T_QA, "两批结构化问答与实际选择",
        ["问题", "选项", "实际选择", "主题"],
        qa_rows, widths=[4.8, 4.2, 3.0, 2.6], font_size=8.5,
    )
    note = body(doc, "完整问答原文见 logs/interactions_20261023_audit02.jsonl。")
    note.alignment = 0


def _editor_note(index: int) -> str:
    return {
        9: "确定自主公开数据、均衡研究对比和暂无硬截止，后续数据与对照按此执行。",
        12: "确定 E 盘独立资源目录和三个实验均衡对比，环境与资源路径据此配置。",
        16: "第一版实验设计与写作计划；后续按用户要求保存为 EXPERIMENT_PLAN.md。",
        19: "视觉部分独立目录与实施计划；最终落盘为 vision/ 布局。",
        35: "首次发现 TF32 导致的分数差异，触发严格 FP32 重跑。",
        36: "严格 FP32 重跑结果；图像最大差 4.1723×10⁻⁷，摘要预测一致。",
        43: "要求把学生交互与子代理交互分开，直接导致附录 C/D 拆分。",
        46: "补回两批结构化问答；确认 62 条代理正文加密，报告保留缺口。",
    }.get(index, "见正文对应章节。")


def _role(value: str) -> str:
    value = str(value or "")
    mapping = {
        "/root": "主",
        "/root/review_experiment_plan": "规划",
        "/root/image_experiment": "图像",
        "/root/video_experiment": "视频",
        "/root/summary_experiment": "摘要",
        "review_experiment_plan": "规划",
        "image_experiment": "图像",
        "video_experiment": "视频",
        "summary_experiment": "摘要",
        "main_agent": "主",
    }
    if value in mapping:
        return mapping[value]
    return value.rsplit("/", 1)[-1][:12]


def add_appendix_d(doc, ctx):
    agents = ctx["agent_records"]
    heading(doc, "附录 D 代理协作记录（不计入学生轮数）", 1)
    body(
        doc,
        "本附录列出 79 条代理协作事件，其中 17 条正文可读，62 条在本地会话日志中只保留加密正文。"
        "这些事件由主代理或子代理发起，不是学生直接输入，也不计入学生交互轮数。"
        "本附录只使用时间、发送方、接收方、类型、模型和可读正文；没有读取密钥，没有尝试解密，也没有把事后概述写成原始指令。",
    )
    add_table(
        doc, T_DSTAT, "代理协作记录统计",
        ["口径", "数量"],
        [
            ["代理协作事件", str(len(agents))],
            ["可读正文", str(sum(1 for r in agents if r.get("body_available")))],
            ["加密正文", str(sum(1 for r in agents if not r.get("body_available")))],
            ["计入学生轮数", "0"],
        ],
        widths=[5.6, 9.0], font_size=9,
    )

    rows = []
    model_short = {"gpt-6-astra": "6-astra", "gpt-6.1-sol": "6.1-sol"}
    for index, record in enumerate(agents, 1):
        rows.append([
            f"D-{index:02d}", _time(record)[11:19],
            f"{_role(record.get('sender', ''))}→{_role(record.get('recipient', ''))}",
            _kind_zh(record.get("kind")),
            model_short.get(record.get("model", ""), record.get("model", "")),
            "可读" if record.get("body_available") else "加密",
        ])
    add_table(
        doc, T_DTABLE, "全部 79 条代理协作记录",
        ["编号", "时间", "方向", "类型", "模型", "正文"],
        rows, widths=[1.8, 2.4, 3.8, 2.2, 2.4, 2.0], font_size=7.5,
        align_center_cols={0, 1, 5}, keep_rows=True,
    )

    heading(doc, "D.2 可读正文", 2)
    readable = [r for r in agents if r.get("body_available")]
    intro = body(
        doc,
        f"17 条正文可读，下面保留前 8 条作为代表；其余 {len(readable) - 8} 条仍列在表 27 的“正文”列为可读的行中，"
        "完整原文保存在 logs/interactions_20261023_audit02.jsonl。",
    )
    intro.alignment = 0
    for record in readable[:8]:
        index = agents.index(record) + 1
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(f"D-{index:02d}  {_short(record.get('text', ''), 40)}")
        _set_run_fonts(run, "宋体", 9, True, None)
        add_box(
            doc, clean_latex(record.get("text", "")).split("\n"),
            fill="FFFFFF", border_color="BFBFBF", font_size=8,
            label=f"{_role(record.get('sender', ''))} → {_role(record.get('recipient', ''))}",
            keep_together=False,
        )

    heading(doc, "D.3 不可读正文清单", 2)
    body(
        doc,
        "62 条事件的正文在会话持久化文件中是加密的。报告只保留编号、时间、方向、类型和模型，"
        "不把空白补成原文，也不根据上下文猜测内容。需要严格完整交互时，这一段就是明确缺口。",
    )
    gap_rows = []
    for index, record in enumerate([r for r in agents if not r.get("body_available")], 1):
        gap_rows.append([
            f"G-{index:02d}", _time(record)[11:19],
            f"{_role(record.get('sender', ''))}→{_role(record.get('recipient', ''))}",
            _kind_zh(record.get("kind")),
        ])
    add_table(
        doc, T_DGAP, "62 条加密正文的元数据",
        ["编号", "时间", "方向", "类型"],
        gap_rows, widths=[1.8, 2.6, 6.2, 4.0], font_size=8,
        align_center_cols={0, 1, 3},
    )
